"""
Synthetic data generator for the BlueWave Logistics ocean freight portfolio project.
All companies, carriers, customers and figures are fictional. Port codes are real UN/LOCODEs.
Run from the repo root:  python scripts/generate_data.py   (writes into data/raw/)
"""
import random, csv, os
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

random.seed(42)
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "raw")
os.makedirs(os.path.join(ROOT, "shipments"), exist_ok=True)
TODAY = date(2026, 10, 24)

# ---------------- Master data ----------------
carriers = [
    # code, canonical name, alliance, on-time probability, name variants seen in raw files
    ("PSL", "Pacific Star Line", "Horizon", 0.78, ["Pacific Star Line", "PACIFIC STAR LINE", "Pacific Star Line Ltd", " pacific star line"]),
    ("OCX", "Oceanic Express", "Horizon", 0.66, ["Oceanic Express", "OCEANIC EXPRESS", "Oceanic Express "]),
    ("JDL", "Jade Dragon Shipping", "Meridian", 0.58, ["Jade Dragon Shipping", "Jade Dragon", "JADE DRAGON SHIPPING"]),
    ("NTC", "Northern Tide Carriers", "Meridian", 0.71, ["Northern Tide Carriers", "Northern Tide", "northern tide carriers"]),
    ("BLM", "Blue Meridian Lines", "Independent", 0.52, ["Blue Meridian Lines", "BLUE MERIDIAN LINES", "Blue Meridian"]),
    ("ATS", "Atlas Trans-Pacific", "Horizon", 0.69, ["Atlas Trans-Pacific", "Atlas TransPacific", "ATLAS TRANS-PACIFIC"]),
    ("CRS", "Coral Sea Container", "Independent", 0.62, ["Coral Sea Container", "Coral Sea Container Co.", "CORAL SEA CONTAINER"]),
]
origins = [("CNSHA","Shanghai","China","Asia"),("CNNGB","Ningbo","China","Asia"),("CNYTN","Yantian","China","Asia"),
           ("CNTAO","Qingdao","China","Asia"),("KRPUS","Busan","South Korea","Asia"),("VNSGN","Ho Chi Minh City","Vietnam","Asia"),
           ("TWKHH","Kaohsiung","Taiwan","Asia")]
dests = [("USLAX","Los Angeles","United States","US West Coast"),("USLGB","Long Beach","United States","US West Coast"),
         ("USOAK","Oakland","United States","US West Coast"),("USSEA","Seattle","United States","US West Coast"),
         ("USNYC","New York/New Jersey","United States","US East Coast"),("USSAV","Savannah","United States","US East Coast"),
         ("USBAL","Baltimore","United States","US East Coast"),("USHOU","Houston","United States","US Gulf Coast")]
origin_w = [30,20,16,10,9,10,5]
dest_w = [26,18,6,7,17,13,6,7]
transit_base = {"US West Coast":(14,19),"US East Coast":(30,38),"US Gulf Coast":(28,35)}
rate_base = {"US West Coast":2300,"US East Coast":3500,"US Gulf Coast":3300}   # USD per FEU, early-2025 baseline

containers = [("20GP","20' Standard",1.0,"Dry",0.55),("40GP","40' Standard",2.0,"Dry",1.0),
              ("40HC","40' High Cube",2.0,"Dry",1.05),("45HC","45' High Cube",2.25,"Dry",1.2),
              ("20RF","20' Reefer",1.0,"Reefer",1.1),("40RF","40' Reefer",2.0,"Reefer",1.9)]
cont_w = [20,22,40,4,5,9]

industries = ["Retail","Consumer Electronics","Furniture","Automotive Parts","Apparel","Home Improvement","Food & Beverage","Industrial Equipment"]
regions = ["East","Central","West"]
managers = {"East":"Linda Chen","Central":"Marcus Reed","West":"Sofia Alvarez"}
name_a = ["Summit","Harbor","Evergreen","Keystone","Redwood","Lakeside","Ironbridge","Silverline","Northfield","Crescent",
          "Granite","Brightway","Oakmont","Riverstone","Pinecrest","Bayview","Highland","Cobalt","Maple Ridge","Sterling",
          "Westgate","Clearwater","Falcon","Meadowbrook","Atlas Point"]
name_b = ["Home Goods","Electronics","Furnishings","Auto Supply","Apparel Co.","Hardware","Foods","Industrial"]
customers = []
for i, a in enumerate(name_a):
    ind_i = i % len(industries)
    reg = regions[i % 3]
    customers.append((f"C{1001+i}", f"{a} {name_b[ind_i]}", industries[ind_i], reg, managers[reg],
                      random.choice(["Strategic","Key","Standard","Standard"])))
cust_w = [random.randint(1,10) * (3 if c[5]=="Strategic" else 2 if c[5]=="Key" else 1) for c in customers]

delay_reasons = ["Port Congestion","Weather","Vessel Rollover","Customs Hold","Equipment Shortage","Schedule Change"]
incoterms = ["FOB","FOB","FOB","CIF","EXW","DDP"]

def season_factor(d):
    return {1:0.95,2:0.70,3:0.90,4:0.95,5:1.0,6:1.05,7:1.15,8:1.25,9:1.25,10:1.15,11:1.0,12:0.9}[d.month]

def rate_index(d):
    months = (d.year-2025)*12 + d.month - 1
    curve = [1.00,0.96,0.92,0.90,0.95,1.10,1.25,1.30,1.18,1.02,0.95,0.97,
             1.00,0.90,0.88,0.93,1.05,1.30,1.45,1.40,1.22,1.10,1.05]
    return curve[min(months, len(curve)-1)]

# ---------------- Shipments ----------------
rows = []
sid = 250001
d, end = date(2025,1,1), date(2026,9,30)
while d <= end:
    n = max(0, int(random.gauss(5.0*season_factor(d), 1.6)))
    for _ in range(n):
        o = random.choices(origins, origin_w)[0]
        de = random.choices(dests, dest_w)[0]
        car = random.choice(carriers)
        ct = random.choices(containers, cont_w)[0]
        cu = random.choices(customers, cust_w)[0]
        qty = random.choice([1,1,1,2,2,3,4,5])
        etd = d + timedelta(days=random.randint(7,21))
        lo, hi = transit_base[de[3]]
        if o[0] == "KRPUS": lo, hi = lo-2, hi-2
        if o[0] == "VNSGN": lo, hi = lo+3, hi+3
        eta = etd + timedelta(days=random.randint(lo, hi))
        dep_delay, reason = 0, ""
        if random.random() < 0.08:
            dep_delay, reason = random.choice([3,7,7,14]), "Vessel Rollover"
        p_on = car[3]
        congested = (de[0] in ("USLAX","USLGB") and eta.month in (8,9,10)) or \
                    (de[0]=="USSAV" and date(2025,7,1)<=eta<=date(2025,10,31)) or \
                    (de[0]=="USNYC" and eta.year==2026 and eta.month in (1,2))
        if congested: p_on -= 0.25
        arr_noise = random.randint(-2,1) if random.random() < p_on else random.randint(2,9)
        if congested and arr_noise > 1 and not reason: reason = "Port Congestion"
        atd = etd + timedelta(days=dep_delay)
        ata = eta + timedelta(days=dep_delay + arr_noise)
        late_days = (ata - eta).days
        if late_days > 1 and not reason:
            reason = random.choices(delay_reasons, [25,15,0,15,15,30])[0]
        if late_days <= 1 and reason != "Vessel Rollover": reason = ""
        status = "Delivered"
        if atd > TODAY: atd_v, ata_v, status, reason = None, None, "Booked", ""
        elif ata > TODAY: atd_v, ata_v, status = atd, None, "In Transit"
        else: atd_v, ata_v = atd, ata
        if status == "In Transit" and reason != "Vessel Rollover": reason = ""
        cost = round(rate_base[de[3]] * rate_index(etd) * random.uniform(0.9,1.12) * ct[4] * qty, 2)
        margin = random.uniform(0.08, 0.22) * (0.8 if cu[5]=="Strategic" else 1.0)
        rows.append(dict(ShipmentID=f"BW{sid}", BookingDate=d, CustomerID=cu[0], CarrierName=random.choice(car[4]),
                         OriginPort=o[0], DestinationPort=de[0], ContainerType=ct[0], ContainerQty=qty,
                         Incoterm=random.choice(incoterms), ETD=etd, ATD=atd_v, ETA=eta, ATA=ata_v,
                         FreightCost_USD=cost, Revenue_USD=round(cost*(1+margin),2), Status=status, DelayReason=reason))
        sid += 1
    d += timedelta(days=1)

# ---------------- Monthly CSVs with intentional data-quality issues ----------------
cols = list(rows[0].keys())
eu_months = {"2025-08","2026-03"}   # exported from a European agent system
dup_ids = {rows[i]["ShipmentID"] for i in random.sample(range(len(rows)), 18)}
by_month = {}
for r in rows:
    by_month.setdefault(r["BookingDate"].strftime("%Y-%m"), []).append(r)

def fdate(v, eu):
    if v is None: return ""
    return v.strftime("%d/%m/%Y") if eu else v.strftime("%Y-%m-%d")
def fnum(v, eu, as_text):
    if eu: return f"{v:.2f}".replace(".", ",")
    return f"${v:,.2f}" if as_text else f"{v:.2f}"

for ym, rs in by_month.items():
    eu = ym in eu_months
    with open(os.path.join(ROOT, "shipments", f"shipments_{ym}.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";" if eu else ",")
        w.writerow(cols)
        for r in rs:
            pc = (lambda s: s.lower()) if random.random() < 0.03 else (lambda s: s)
            line = [r["ShipmentID"], fdate(r["BookingDate"], eu), r["CustomerID"], r["CarrierName"],
                    pc(r["OriginPort"]), pc(r["DestinationPort"]), r["ContainerType"], r["ContainerQty"], r["Incoterm"],
                    fdate(r["ETD"], eu), fdate(r["ATD"], eu), fdate(r["ETA"], eu), fdate(r["ATA"], eu),
                    fnum(r["FreightCost_USD"], eu, random.random() < 0.04), fnum(r["Revenue_USD"], eu, False),
                    r["Status"], r["DelayReason"]]
            w.writerow(line)
            if r["ShipmentID"] in dup_ids: w.writerow(line)
        if random.random() < 0.3:
            w.writerow([""]*len(cols)); w.writerow([""]*len(cols))

# ---------------- master_data.xlsx ----------------
F, FB, FT = Font(name="Arial"), Font(name="Arial", bold=True), Font(name="Arial", bold=True, size=13)
HDR = PatternFill("solid", fgColor="D9E1F2")
def sheet(ws, header, data, title_rows=None):
    r0 = 1
    if title_rows:
        for i, t in enumerate(title_rows):
            ws.cell(row=1+i, column=1, value=t).font = FT if i == 0 else F
        r0 = len(title_rows) + 2
    for j, h in enumerate(header, 1):
        c = ws.cell(row=r0, column=j, value=h); c.font, c.fill = FB, HDR
        ws.column_dimensions[c.column_letter].width = max(14, len(h)+6)
    for i, row in enumerate(data, r0+1):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=v).font = F

wb = Workbook()
ws = wb.active; ws.title = "Carriers"
sheet(ws, ["CarrierCode","CarrierName","Alliance"], [(c[0],c[1],c[2]) for c in carriers])
sheet(wb.create_sheet("Ports"), ["PortCode","PortName","Country","Region","PortRole"],
      [(p[0],p[1],p[2],p[3],"Origin") for p in origins] + [(p[0],p[1],p[2],p[3],"Destination") for p in dests])
sheet(wb.create_sheet("Customers"), ["CustomerID","CustomerName","Industry","SalesRegion","AccountManager","Tier"], customers,
      title_rows=["BlueWave Logistics - Customer Master", "Exported from CRM on 2026-10-01 | Internal use only"])
sheet(wb.create_sheet("ContainerTypes"), ["ContainerType","Description","TEU_Factor","Category"], [c[:4] for c in containers])
wb.save(os.path.join(ROOT, "master_data.xlsx"))

# ---------------- budget_2026.xlsx (wide format: months as columns) ----------------
teu25 = {}
dreg = {x[0]: x[3] for x in dests}; tfac = {c[0]: c[2] for c in containers}
for r in rows:
    if r["BookingDate"].year == 2025:
        k = (dreg[r["DestinationPort"]], r["BookingDate"].month)
        teu25[k] = teu25.get(k, 0) + tfac[r["ContainerType"]] * r["ContainerQty"]
regs = ["US West Coast","US East Coast","US Gulf Coast"]
per_teu = {"US West Coast":1500,"US East Coast":2150,"US Gulf Coast":2050}
data = [[g, "TEU"] + [round(teu25.get((g,m), 0)*1.08) for m in range(1,13)] for g in regs] + \
       [[g, "Revenue_USD"] + [round(teu25.get((g,m), 0)*1.08*per_teu[g], -2) for m in range(1,13)] for g in regs]
wb = Workbook(); ws = wb.active; ws.title = "Budget_2026"
sheet(ws, ["DestinationRegion","Metric","Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"], data)
for row in ws.iter_rows(min_row=2, min_col=3):
    for c in row: c.number_format = "#,##0"
wb.save(os.path.join(ROOT, "budget_2026.xlsx"))

print(f"Shipments: {len(rows)} | monthly files: {len(by_month)} | duplicate rows injected: {len(dup_ids)}")
